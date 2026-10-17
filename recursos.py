"""Conferência e exportação dos dados da caravana."""
from io import BytesIO
import re
import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from geracao import CONVENIO, CONTATO, validar_opcoes, separar_telefone


def validar_caravana(horario='', ddd='', telefone=''):
    horario, ddd, telefone = (v.strip() for v in (horario, ddd, telefone))
    if horario:
        if not re.fullmatch(r'\d{1,2}:\d{2}', horario):
            raise ValueError('Horário de retorno: use HH:MM, por exemplo 18:00, ou deixe em branco.')
        hora, minuto = map(int, horario.split(':'))
        if hora > 23 or minuto > 59:
            raise ValueError('Horário de retorno: informe um horário entre 00:00 e 23:59.')
        horario = f'{hora:02d}:{minuto:02d}'
    if ddd and not re.fullmatch(r'[1-9][0-9]', ddd):
        raise ValueError('DDD do responsável: informe dois dígitos, por exemplo 11.')
    if telefone:
        extraido, numero = separar_telefone(telefone)
        if extraido:
            if ddd and ddd != extraido:
                raise ValueError('O DDD informado difere do DDD no telefone do responsável. Confira os dois campos.')
            ddd, telefone = extraido, numero
        else:
            if not re.fullmatch(r'[0-9\s().-]+', telefone):
                raise ValueError('Telefone do responsável: informe um número brasileiro sem ramal.')
            telefone = re.sub(r'\D', '', telefone)
            if len(telefone) not in (8, 9):
                raise ValueError('Telefone do responsável: use 8 ou 9 dígitos ou informe o número completo com DDD.')
        if not ddd:
            raise ValueError('Informe o DDD do responsável junto ao telefone ou no campo DDD.')
    elif ddd:
        raise ValueError('Informe o telefone do responsável ou deixe DDD e telefone em branco.')
    return horario, ddd, telefone


def exportar_planilha(df):
    """Exporta os participantes atuais como texto, inclusive as escolhas do editor."""
    df = validar_opcoes(df)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Participantes'
    for linha, valores in enumerate([list(df.columns), *df.fillna('').values.tolist()], 1):
        for coluna, valor in enumerate(valores, 1):
            if linha > 1 and df.columns[coluna - 1] in (CONVENIO, CONTATO):
                valor = str(valor).upper()
            cell = ws.cell(linha, coluna, str(valor))
            # Conteúdo de usuário é texto literal, inclusive se começar por =.
            cell.data_type = 's'
            cell.number_format = '@'
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='215E70')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for i, nome in enumerate(df.columns, 1):
        letra = get_column_letter(i)
        ws.column_dimensions[letra].width = 42 if 'Nome' in nome else 28
        if nome in (CONVENIO, CONTATO, 'Trajeto'):
            opcoes = 'IDA E VOLTA,SOMENTE IDA,SOMENTE VOLTA' if nome == 'Trajeto' else 'Sim,Não' if nome == CONVENIO else 'Familiar,Amigo'
            validacao = DataValidation(type='list', formula1=f'"{opcoes}"', allow_blank=True)
            validacao.errorTitle = 'Selecione uma opção'
            validacao.error = 'Use uma das opções da lista ou deixe em branco.'
            validacao.showErrorMessage = True
            ws.add_data_validation(validacao)
            validacao.add(f'{letra}2:{letra}{max(2, len(df) + 1)}')
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def renderizar_pagina(pdf, pagina):
    """Renderiza somente a página solicitada, sem salvar dados no servidor."""
    with pymupdf.open(stream=pdf, filetype='pdf') as doc:
        if not 0 <= pagina < len(doc):
            raise ValueError('Página fora do intervalo do PDF.')
        page = doc[pagina]
        escala = min(1.5, 1100 / page.rect.width)
        return page.get_pixmap(matrix=pymupdf.Matrix(escala, escala), alpha=False).tobytes('png')
